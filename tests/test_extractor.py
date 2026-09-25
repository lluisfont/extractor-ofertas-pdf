"""Prueba de extremo a extremo con un folleto sintético. Ejecutar: python -m pytest tests -q"""
from pathlib import Path

import pytest

from extractor_ofertas.config import Config
from extractor_ofertas.exportar import exportar
from extractor_ofertas.procesador import procesar_pdf

from generar_folleto_prueba import generar


@pytest.fixture(scope="module")
def resultado(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("folleto")
    pdf = carpeta / "folleto_prueba.pdf"
    esperadas = generar(pdf)
    cfg = Config()
    cfg.ocr_activo = False
    res = procesar_pdf(pdf, cfg)
    return res, esperadas, carpeta


def test_todas_las_ofertas(resultado):
    res, esperadas, _ = resultado
    assert sorted(o.campos["precio_oferta"] for o in res.ofertas) == sorted(e["precio"] for e in esperadas)


def test_campos_por_oferta(resultado):
    res, esperadas, _ = resultado
    por_precio = {o.campos["precio_oferta"]: o for o in res.ofertas}
    for e in esperadas:
        o = por_precio[e["precio"]]
        c = o.campos
        assert e["producto"] in c["producto"], (e, c)
        assert c.get("marca", "") == e["marca"], (e, c)
        assert c.get("precio_anterior") == e["anterior"], (e, c)
        if e["pu"]:
            assert c.get("precio_unitario") == e["pu"][0] and c["unidad_precio_unitario"] == "€/" + e["pu"][1], (e, c)
        if e["promo"] and "tarjeta" not in e["promo"]:
            assert e["promo"] in c["promocion"], (e, c)
        assert o.pagina == e["pagina"]


def test_auditoria(resultado):
    res, _, _ = resultado
    assert res.vigencia == "del 1 al 14 de octubre de 2026"
    assert not res.huerfanos
    # La página 3 es una imagen: sin OCR debe quedar marcada como ERROR
    assert res.paginas[2].estado == "ERROR"
    assert res.paginas[0].estado == res.paginas[1].estado == "OK"
    assert not [i for i in res.incidencias if i.tipo == "Conteo de control"]


def test_excel(resultado):
    res, _, carpeta = resultado
    destino = exportar(res, carpeta)
    from openpyxl import load_workbook
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Resumen", "Ofertas", "Auditoria_paginas", "Incidencias", "Texto_bruto"]
    assert wb["Ofertas"].max_row == len(res.ofertas) + 1
