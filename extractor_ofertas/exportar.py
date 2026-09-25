"""Exportación del resultado a Excel con hojas de ofertas, auditoría e incidencias."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from .procesador import Resultado

EURO = '#,##0.00 "€"'
CAB = PatternFill("solid", fgColor="1F4E78")
AVISO = PatternFill("solid", fgColor="FFF2CC")
ERROR = PatternFill("solid", fgColor="F8CBAD")
OK = PatternFill("solid", fgColor="E2EFDA")

COLUMNAS_OFERTAS = [
    ("ID", 6, None), ("Página", 8, None), ("Sección", 18, None), ("Producto", 45, None), ("Marca", 18, None),
    ("Formato", 18, None), ("Precio oferta", 12, EURO), ("Nota precio oferta", 24, None),
    ("Precio normal / anterior", 13, EURO), ("Descuento %", 11, "0.0"), ("Promoción", 20, None),
    ("Precio total lote", 12, EURO), ("Uds. lote", 9, None), ("Precio 2ª unidad", 12, EURO), ("Cupón €", 10, EURO),
    ("Precio unitario", 12, EURO), ("Unidad PU", 10, None), ("PU normal", 11, EURO), ("Condiciones", 24, None),
    ("Vigencia oferta", 18, None), ("Confianza", 10, None), ("Origen precio", 12, None), ("Posición (x, y)", 14, None),
    ("Alertas", 40, None), ("Texto completo del bloque", 70, None),
]


def _hoja(wb, titulo, columnas):
    ws = wb.create_sheet(titulo)
    for j, (nombre, ancho, _) in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.freeze_panes = "A2"
    return ws


def _filas(ws, columnas, filas, relleno=None):
    for i, fila in enumerate(filas, start=2):
        for j, valor in enumerate(fila, start=1):
            if isinstance(valor, str):
                valor = ILLEGAL_CHARACTERS_RE.sub("", valor) or None
            c = ws.cell(row=i, column=j, value=valor)
            fmt = columnas[j - 1][2]
            if fmt:
                c.number_format = fmt
        if relleno:
            color = relleno(fila)
            if color:
                for j in range(1, len(columnas) + 1):
                    ws.cell(row=i, column=j).fill = color
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"


def exportar(res: Resultado, carpeta: Path) -> Path:
    carpeta.mkdir(parents=True, exist_ok=True)
    destino = carpeta / f"{res.archivo.stem}_ofertas.xlsx"
    if destino.exists():
        destino = carpeta / f"{res.archivo.stem}_ofertas_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)

    # --- Ofertas ---
    ws = _hoja(wb, "Ofertas", COLUMNAS_OFERTAS)
    filas = []
    for of in res.ofertas:
        c = of.campos
        filas.append([
            of.id, of.pagina, of.seccion, c.get("producto"), c.get("marca"), c.get("formato"),
            c.get("precio_oferta"), c.get("nota_precio"), c.get("precio_anterior"), c.get("descuento_pct"),
            c.get("promocion"), c.get("precio_total_lote"), c.get("unidades_lote"), c.get("precio_segunda_unidad"),
            c.get("cupon"), c.get("precio_unitario"), c.get("unidad_precio_unitario"), c.get("precio_unitario_normal"),
            c.get("condiciones"), c.get("vigencia"), of.precio.confianza, of.precio.motor,
            (f"{of.precio.caja.x0:.0f}, {of.precio.caja.y0:.0f}" if of.precio.caja.x0 or of.precio.caja.y0 else ""), "; ".join(of.alertas), c.get("texto_completo"),
        ])
    _filas(ws, COLUMNAS_OFERTAS, filas, relleno=lambda f: AVISO if f[23] else None)

    # --- Resumen ---
    rs = wb.create_sheet("Resumen", 0)
    estado = res.estado
    datos = [
        ("Archivo", res.archivo.name),
        ("Procesado", res.inicio.strftime("%d/%m/%Y %H:%M:%S")),
        ("Vigencia detectada", res.vigencia or "(no detectada)"),
        ("Páginas", len(res.paginas)),
        ("Ofertas extraídas", len(res.ofertas)),
        ("Precios € en texto bruto (PyMuPDF)", sum(p.bruto_pymupdf for p in res.paginas)),
        ("Precios € en texto bruto (pdfplumber)", sum(p.bruto_pdfplumber for p in res.paginas)),
        ("Precios detectados (todos los tipos)", sum(p.precios_detectados for p in res.paginas)),
        ("Precios recuperados por triangulación", sum(p.recuperados_triangulacion for p in res.paginas)),
        ("Precios huérfanos (sin oferta)", len(res.huerfanos)),
        ("Ofertas sin precio (solo sello promo)", sum(1 for o in res.ofertas if o.precio.tipo == "sello")),
        ("Ofertas con alertas", sum(1 for o in res.ofertas if o.alertas)),
        ("Páginas leídas con OCR", sum(1 for p in res.paginas if p.metodo == "OCR")),
        ("OCR disponible", "Sí" if res.ocr_disponible else "No"),
        ("Avisos / errores", f"{sum(i.gravedad == 'AVISO' for i in res.incidencias)} / "
                             f"{sum(i.gravedad == 'ERROR' for i in res.incidencias)}"),
        ("ESTADO DE VERIFICACIÓN", estado),
    ]
    for i, (k, v) in enumerate(datos, start=1):
        rs.cell(row=i, column=1, value=k).font = Font(bold=True)
        rs.cell(row=i, column=2, value=v)
    ultima = rs.cell(row=len(datos), column=2)
    ultima.font = Font(bold=True)
    ultima.fill = ERROR if "errores" in estado else (AVISO if "AVISOS" in estado else OK)
    rs.column_dimensions["A"].width = 40
    rs.column_dimensions["B"].width = 60

    # --- Auditoría por página ---
    cols = [("Página", 8, None), ("Método", 9, None), ("Palabras", 10, None), ("€ bruto PyMuPDF", 14, None),
            ("€ bruto pdfplumber", 15, None), ("Precios detectados", 14, None), ("Principales", 11, None),
            ("Secundarios", 11, None), ("Recuperados triangulación", 15, None), ("Huérfanos", 10, None),
            ("Ofertas", 9, None), ("% imagen", 9, "0%"), ("Estado", 10, None)]
    wa = _hoja(wb, "Auditoria_paginas", cols)
    filas = [[p.numero, p.metodo, p.palabras, p.bruto_pymupdf, p.bruto_pdfplumber, p.precios_detectados,
              p.precios_principales, p.precios_secundarios, p.recuperados_triangulacion, p.huerfanos, p.ofertas,
              p.cobertura_imagen, p.estado] for p in res.paginas]
    _filas(wa, cols, filas, relleno=lambda f: ERROR if f[12] == "ERROR" else (AVISO if f[12] == "REVISAR" else None))

    # --- Incidencias ---
    cols = [("Página", 8, None), ("Gravedad", 10, None), ("Tipo", 22, None), ("Detalle", 110, None)]
    wi = _hoja(wb, "Incidencias", cols)
    orden = {"ERROR": 0, "AVISO": 1, "INFO": 2}
    filas = [[i.pagina, i.gravedad, i.tipo, i.detalle]
             for i in sorted(res.incidencias, key=lambda x: (orden[x.gravedad], x.pagina or 0))]
    _filas(wi, cols, filas, relleno=lambda f: ERROR if f[1] == "ERROR" else (AVISO if f[1] == "AVISO" else None))

    # --- Texto bruto (trazabilidad de cada línea) ---
    cols = [("Página", 8, None), ("Oferta ID", 9, None), ("x", 7, "0"), ("y", 7, "0"), ("Tamaño letra", 11, "0.0"),
            ("Motor", 11, None), ("Texto de la línea", 90, None)]
    wt = _hoja(wb, "Texto_bruto", cols)
    filas = [[n, oid, l.x0, l.y0, l.tamano, l.palabras[0].motor, l.texto] for n, l, oid in res.lineas_brutas]
    _filas(wt, cols, filas)

    wb.save(destino)
    return destino
